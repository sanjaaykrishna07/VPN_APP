import os
from pptx import Presentation
from pptx.util import Inches

def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    slide_titles = [
        "Title Slide: QuantumSafe VPN",
        "Motivation & Threat Landscape",
        "Problem Statement & Research Gaps",
        "Project Objectives",
        "System Architecture & Network Topology",
        "Experimental Results & Security Validation",
        "Thank You"
    ]

    for i in range(1, 8):
        slide_img = f"d:/VPN_APP/slides/slide_{i}.png"
        if not os.path.exists(slide_img):
            print(f"Warning: {slide_img} not found!")
            continue
        
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(slide_img, Inches(0), Inches(0), width=Inches(13.333), height=Inches(7.5))
        
        # Add slide title to notes for indexing / accessibility
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = f"Slide {i}: {slide_titles[i-1]}"

    output_path = r"d:\VPN_APP\QuantumSafeVPN_Presentation.pptx"
    prs.save(output_path)
    print(f"Successfully generated {output_path} with 7 slides!")

if __name__ == "__main__":
    build_pptx()
